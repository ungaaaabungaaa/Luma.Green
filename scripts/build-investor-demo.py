#!/usr/bin/env python3
"""Build a draft tour; final publication requires fresh captures and page review."""
from __future__ import annotations
import argparse
import hashlib
import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

from document_links import add_document_hyperlink

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'docs/investor-demo/tour.json'
OUTPUT = ROOT / 'output/docx/luma-green-investor-demo.docx'
RECORD = ROOT / 'docs/investor-demo/build.json'


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def validate(data: dict, final: bool) -> list[Path]:
    if not 10 <= len(data['pages']) <= 14:
        raise ValueError('The tour must have 10 to 14 planned pages.')
    if (data['accounts_per_environment'], data['persona_templates'], data['accounts_per_template'], data['admin_accounts']) != (140, 28, 5, 0):
        raise ValueError('The declared investor roster differs from the agreed scope.')
    images = []
    for index, page in enumerate(data['pages']):
        for key in ['title', 'role', 'intro', 'result', 'limit']:
            if not isinstance(page.get(key), str) or not page[key].strip():
                raise ValueError('Every page needs a title, role, purpose, result and limit.')
        shot = page.get('screenshot')
        if not shot:
            if final and index < 12:
                raise ValueError('Final tour requires new captures for each demonstration page.')
            continue
        path = (SOURCE.parent / shot['path']).resolve()
        if not path.is_relative_to((SOURCE.parent / 'screenshots').resolve()) or not path.is_file():
            raise ValueError('Capture must be in docs/investor-demo/screenshots.')
        if sha(path) != shot['sha256']:
            raise ValueError('Capture bytes differ from their reviewed evidence.')
        if shot.get('reviewed') is not True or shot.get('actual_browser_capture') is not True:
            raise ValueError('Capture requires explicit original-image browser review.')
        if shot.get('environment') not in {'production_demo', 'development_demo'}:
            raise ValueError('Capture needs an explicit demo environment.')
        for key in ['caption', 'release', 'dataset', 'captured_at']:
            if not shot.get(key):
                raise ValueError('Capture provenance is incomplete.')
        with Image.open(path) as image:
            if image.width < 700:
                raise ValueError('Use a readable screenshot at least 700 pixels wide.')
        images.append(path)
    if final and not data.get('development_url'):
        raise ValueError('Confirm the development frontend address before final publication.')
    return images


def configure(document: Document) -> None:
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.6)
    section.left_margin = section.right_margin = Inches(.7)
    for name in ['Normal', 'Title', 'Heading 1', 'Caption', 'List Number']:
        style = document.styles[name]
        style.font.name = 'DejaVu Sans'
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.widow_control = True
    for border in list(document.styles.element.xpath('.//w:pBdr')):
        border.getparent().remove(border)
    document.styles['Caption'].font.bold = False
    document.styles['Normal'].font.size = Pt(11)
    document.styles['Normal'].paragraph_format.line_spacing = 1.08
    document.styles['List Number'].font.size = Pt(11)
    document.styles['Title'].font.size = Pt(25)
    document.styles['Heading 1'].font.size = Pt(20)
    document.styles['Caption'].font.size = Pt(9)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('Luma Green investor walkthrough · ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    document.core_properties.title = 'Luma Green investor walkthrough'
    document.core_properties.author = 'Luma.Green'
    document.core_properties.subject = 'Imported demonstration records and workflow tour'


def paragraph(document: Document, label: str, text: str) -> None:
    p = document.add_paragraph()
    p.add_run(label + ' ').bold = True
    p.add_run(text)


def build(final: bool) -> None:
    data = json.loads(SOURCE.read_text())
    images = validate(data, final)
    document = Document()
    configure(document)
    for index, page in enumerate(data['pages']):
        if index:
            document.add_page_break()
        if index == 0:
            document.add_paragraph(data['title'], 'Title')
            document.add_paragraph(data['edition'], 'Caption')
        document.add_paragraph(f'{index + 1:02d}  {page["title"]}', 'Heading 1')
        paragraph(document, 'Role', page['role'])
        document.add_paragraph(page['intro'])
        for number, step in enumerate(page['steps'], 1):
            document.add_paragraph(f'{number}. {step}')
        shot = page.get('screenshot')
        if shot:
            path = (SOURCE.parent / shot['path']).resolve()
            with Image.open(path) as image:
                width, height = image.size
            scale = min(7.1 / width, (3.6 if index == 0 else 4.3) / height)
            document.add_picture(str(path), width=Inches(width * scale), height=Inches(height * scale))
            document.add_paragraph(shot['caption'], 'Caption')
        elif index < 12:
            document.add_paragraph('Screenshot pending new demo capture. No earlier fixture is presented as this environment.', 'Caption')
        paragraph(document, 'Check', page['result'])
        paragraph(document, 'Limit', page['limit'])
        if index == len(data['pages']) - 1:
            for name, note, url in data['research']:
                p = document.add_paragraph()
                add_document_hyperlink(p, name, url, bold=True)
                p.add_run(' — ' + note)
            document.add_paragraph('Official sources checked 10 October 2026. Scenario mappings are interpretation, not evidence of affiliation.', 'Caption')
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    inputs = [SOURCE, Path(__file__).resolve(), ROOT / 'scripts/document_links.py', *images]
    record = dict(status='render_and_review_pending', publication_mode='final_candidate' if final else 'draft_pending_new_captures', planned_pages=len(data['pages']), image_count=len(images), docx=str(OUTPUT.relative_to(ROOT)), docx_sha256=sha(OUTPUT), inputs={str(p.relative_to(ROOT)): sha(p) for p in inputs}, visual_review={'status': 'pending'})
    RECORD.write_text(json.dumps(record, indent=2) + '\n')
    print(f'Built {len(data["pages"])} planned pages with {len(images)} reviewed screenshots. Render review is pending.')


def record_review(pages: int) -> None:
    record = json.loads(RECORD.read_text())
    if not 10 <= pages <= 14:
        raise ValueError('Rendered tour must contain 10 to 14 pages.')
    if sha(OUTPUT) != record['docx_sha256'] or any(sha(ROOT / name) != value for name, value in record['inputs'].items()):
        raise ValueError('Inputs changed after the build. Rebuild and inspect again.')
    record['visual_review'] = {'status': 'passed', 'pages': pages, 'docx_sha256': sha(OUTPUT)}
    record['status'] = 'reviewed_draft' if record['publication_mode'].startswith('draft') else 'reviewed_final_candidate'
    RECORD.write_text(json.dumps(record, indent=2) + '\n')
    print('Recorded the supplied whole-document visual review. Draft status is preserved.')


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--final', action='store_true')
    parser.add_argument('--check', action='store_true')
    parser.add_argument('--record-review', type=int)
    args = parser.parse_args()
    if args.record_review is not None:
        record_review(args.record_review)
    elif args.check:
        validate(json.loads(SOURCE.read_text()), args.final)
        print('Tour source is valid. This does not establish capture, render or provider readiness.')
    else:
        build(args.final)
