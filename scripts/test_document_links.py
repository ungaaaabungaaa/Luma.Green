"""Offline regression tests for portable Word links. Run with the docs Python runtime."""
import importlib.util
import tempfile
import unittest
from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZipFile

from docx import Document
from document_links import resolve_document_link

ROOT = Path(__file__).resolve().parents[1]


def builder(name):
    spec = importlib.util.spec_from_file_location(name, ROOT / "scripts" / name)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


class DocumentLinkTests(unittest.TestCase):
    def test_relative_repository_target_and_anchor_are_portable(self):
        result = resolve_document_link("../design/designer-system.md#current-ui-contract", ROOT / "docs/testing/launch-2026-10-10.md", ROOT)
        self.assertEqual(result, "https://github.com/ungaaaabungaaa/Luma.Green/blob/main/docs/design/designer-system.md#current-ui-contract")

    def test_external_link_remains_exact(self):
        url = "https://docs.convex.dev/production/hosting/vercel?example=1#setup"
        self.assertEqual(resolve_document_link(url, ROOT / "docs/user-guide/guide.md", ROOT), url)

    def test_local_unsafe_and_missing_targets_are_rejected(self):
        for target in ["file:///private/tmp/secret", "/Users/example/secret", "../../../outside.md", "missing-file.md", "javascript:alert(1)", "//example.com/file"]:
            with self.subTest(target=target), self.assertRaises(ValueError):
                resolve_document_link(target, ROOT / "docs/user-guide/guide.md", ROOT)

    def test_both_builders_keep_clickable_links_after_word_roundtrip(self):
        guide = builder("build-user-guide.py")
        planning = builder("build-planning-docs.py")
        doc = Document()
        guide.inline(doc.add_paragraph(), "Read [Setup](../operations/launch-checklist.md).")
        planning.add_inline(doc.add_paragraph(), "Read [Setup](../operations/launch-checklist.md).", (ROOT / "docs/testing/launch-2026-10-10.md", ROOT))
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "links.docx"
            doc.save(output)
            reopened = Document(output)
            expected = "https://github.com/ungaaaabungaaa/Luma.Green/blob/main/docs/operations/launch-checklist.md"
            self.assertIn("Setup (" + expected + ")", reopened.paragraphs[0].text)
            self.assertEqual(reopened.paragraphs[1].text, "Read Setup.")
            for paragraph in reopened.paragraphs:
                self.assertEqual(paragraph.hyperlinks[0].url, expected)
            with ZipFile(output) as archive:
                relationships = ElementTree.fromstring(archive.read("word/_rels/document.xml.rels"))
                targets = [item.attrib["Target"] for item in relationships if item.attrib["Type"].endswith("/hyperlink")]
                self.assertEqual(targets, [expected])


if __name__ == "__main__":
    unittest.main()
