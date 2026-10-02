#!/usr/bin/env python3
"""Build the editable team review pack and verify its reviewed inputs.

Use the bundled Documents runtime. Run the artifact marker before authoring.
After rendering and inspecting every page, use --record-review N, then --check.
"""
from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import re
import zipfile
from pathlib import Path
from urllib.parse import urlparse
from xml.etree import ElementTree

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.shared import Inches, Pt
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
SOURCES = [
    ROOT / 'docs/team-review/review.md',
    ROOT / 'docs/testing/team-end-to-end-manual.md',
    ROOT / 'docs/product/six-month-execution-plan.md',
    ROOT / 'docs/operations/india-entity-trademark-and-legal.md',
]
OUTPUT = ROOT / 'output/docx/luma-green-team-review.docx'
RECORD = ROOT / 'docs/team-review/build.json'
MANIFEST = ROOT / 'docs/team-review/screenshots/manifest.json'
SHARED_BUILDER = ROOT / 'scripts/build-user-guide.py'
BODY_WIDTH = 7.0
IMAGE = re.compile(r'!\[([^\]]+)\]\(([^)]+)\)')
LINK = re.compile(r'\[([^\]]+)\]\(([^)]+)\)')
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}

spec = importlib.util.spec_from_file_location('luma_user_guide_builder', SHARED_BUILDER)
if spec is None or spec.loader is None:
    raise RuntimeError('Cannot load the shared Word style helpers.')
shared = importlib.util.module_from_spec(spec)
spec.loader.exec_module(shared)


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def plain(value: str) -> str:
    value = LINK.sub(lambda match: match[1], value)
    return value.replace('**', '').replace('`', '').replace('\\|', '|')


def normalized(value: str) -> str:
    return ' '.join(value.split())


def image_path(source: Path, relative: str) -> Path:
    if urlparse(relative).scheme:
        raise ValueError('Screenshots must be local browser captures, not remote URLs.')
    path = (source.parent / relative).resolve()
    if not path.is_relative_to(ROOT):
        raise ValueError(f'Screenshot must be inside the repository: {relative}')
    if not path.is_file():
        raise ValueError(f'Missing screenshot: {path}')
    return path


def special(line: str) -> bool:
    return bool(re.match(r'^(?:#{1,6} |\||!\[|<!--|```|---$|\d+[.)] |[-*] )', line))


def blocks(source: Path) -> list[tuple[str, object]]:
    """Parse the small, explicit Markdown subset used by the four sources."""
    lines = source.read_text(encoding='utf-8').splitlines()
    result: list[tuple[str, object]] = []
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        index += 1
        if not line or line == '---':
            continue
        if line == '<!-- pair -->':
            pair = []
            while index < len(lines) and lines[index].strip() != '<!-- /pair -->':
                item = lines[index].strip()
                index += 1
                if not item:
                    continue
                match = IMAGE.fullmatch(item)
                if match is None:
                    raise ValueError(f'{source}: a pair can contain only two Markdown image lines.')
                pair.append((match[1], image_path(source, match[2])))
            if index == len(lines) or len(pair) != 2:
                raise ValueError(f'{source}: each pair needs two images and a closing marker.')
            index += 1
            result.append(('pair', pair))
        elif line == '<!-- /pair -->':
            raise ValueError(f'{source}: unmatched closing pair marker.')
        elif line.startswith('<!--'):
            if not line.endswith('-->'):
                raise ValueError(f'{source}: use single-line comments.')
        elif line.startswith('```'):
            code = []
            while index < len(lines) and not lines[index].strip().startswith('```'):
                code.append(lines[index])
                index += 1
            if index == len(lines):
                raise ValueError(f'{source}: unclosed code fence.')
            index += 1
            result.append(('code', '\n'.join(code)))
        elif match := IMAGE.fullmatch(line):
            result.append(('image', (match[1], image_path(source, match[2]))))
        elif match := re.match(r'^(#{1,6}) (.+)$', line):
            result.append(('heading', (len(match[1]), plain(match[2]))))
        elif line.startswith('|'):
            rows = []
            while True:
                cells = [cell.strip().replace('\\|', '|') for cell in re.split(r'(?<!\\)\|', line.strip('|'))]
                if not all(re.fullmatch(r':?-+:?', cell) for cell in cells):
                    rows.append(cells)
                if index == len(lines) or not lines[index].strip().startswith('|'):
                    break
                line = lines[index].strip()
                index += 1
            if len(rows) < 2 or any(len(row) != len(rows[0]) for row in rows):
                raise ValueError(f'{source}: table needs a header and consistent row widths.')
            result.append(('table', rows))
        else:
            is_list = bool(re.match(r'^(?:\d+[.)]|[-*]) ', line))
            parts = [line]
            while index < len(lines) and lines[index].strip() and not special(lines[index].strip()):
                parts.append(lines[index].strip())
                index += 1
            result.append(('list' if is_list else 'paragraph', ' '.join(parts)))
    return result


def image_uses() -> list[tuple[str, Path]]:
    result = []
    for source in SOURCES:
        for kind, value in blocks(source):
            if kind == 'pair':
                result.extend(value)
            elif kind == 'image':
                result.append(value)
    return result


def inputs() -> dict[str, str]:
    paths = SOURCES + [
        Path(__file__).resolve(), SHARED_BUILDER,
        ROOT / 'scripts/user-guide-requirements.txt',
        ROOT / 'docs/team-review/README.md', MANIFEST,
        ROOT / 'scripts/demo-seed/manifest.json',
        ROOT / 'scripts/demo-seed/coverage.json',
    ] + [path for _, path in image_uses()]
    json.loads(MANIFEST.read_text(encoding='utf-8'))
    return {str(path.relative_to(ROOT)): digest(path) for path in sorted(set(paths))}


def add_inline(paragraph, value: str, bold: bool = False) -> None:
    """Reuse shared runs; print readable labels and retain external hyperlinks."""
    offset = 0
    for match in LINK.finditer(value):
        shared.inline(paragraph, value[offset:match.start()], bold=bold)
        label, url = match.groups()
        if urlparse(url).scheme in {'https', 'http', 'mailto'}:
            link = OxmlElement('w:hyperlink')
            link.set(qn('r:id'), paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
            run = OxmlElement('w:r')
            properties = OxmlElement('w:rPr')
            color = OxmlElement('w:color')
            color.set(qn('w:val'), '000000')
            properties.append(color)
            underline = OxmlElement('w:u')
            underline.set(qn('w:val'), 'single')
            properties.append(underline)
            if bold:
                properties.append(OxmlElement('w:b'))
            run.append(properties)
            text = OxmlElement('w:t')
            text.text = label
            run.append(text)
            link.append(run)
            paragraph._p.append(link)
        else:
            shared.inline(paragraph, label, bold=bold)
        offset = match.end()
    shared.inline(paragraph, value[offset:], bold=bold)


def configure(document) -> None:
    shared.configure(document)
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.left_margin = section.right_margin = Inches(.75)
    normal = document.styles['Normal']
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.1
    normal.paragraph_format.space_after = Pt(6)
    for name, size in [('Title', 28), ('Heading 1', 18), ('Heading 2', 13), ('Heading 3', 11.5)]:
        document.styles[name].font.size = Pt(size)
    document.styles['Caption'].font.size = Pt(9)
    document.core_properties.title = 'Luma Green product review and team test pack'
    document.core_properties.subject = 'Product evidence, role testing, proposed roadmap and legal preparation'


def table_style(table, widths: list[float]) -> None:
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    properties = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement(f'w:{edge}')
        for key, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
            border.set(qn(f'w:{key}'), value)
        borders.append(border)
    properties.append(borders)
    margins = OxmlElement('w:tblCellMar')
    for edge in ['top', 'left', 'bottom', 'right']:
        margin = OxmlElement(f'w:{edge}')
        margin.set(qn('w:w'), '110')
        margin.set(qn('w:type'), 'dxa')
        margins.append(margin)
    properties.append(margins)
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(document, rows: list[list[str]]) -> None:
    if len(rows[0]) >= 5:
        # Wide operational records stay editable and readable at normal page size.
        for row in rows[1:]:
            document.add_heading(shared.heading_text(plain(row[0])), level=3)
            for field_index, (label, value) in enumerate(zip(rows[0], row)):
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.space_after = Pt(4)
                paragraph.paragraph_format.keep_with_next = field_index < len(row) - 1
                add_inline(paragraph, f'**{plain(label)}:** {value}')
        return
    count = len(rows[0])
    proportions = {2: [.28, .72], 3: [.23, .33, .44], 4: [.18, .24, .28, .30]}.get(count, [1])
    widths = [BODY_WIDTH * proportion for proportion in proportions]
    table = document.add_table(rows=len(rows), cols=count)
    table_style(table, widths)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for row_index, (row, values) in enumerate(zip(table.rows, rows)):
        for cell, value in zip(row.cells, values):
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.line_spacing = 1.05
            add_inline(paragraph, value, bold=row_index == 0)
            for run in paragraph.runs:
                run.font.size = Pt(9.5)
            if row_index == 0:
                shade = OxmlElement('w:shd')
                shade.set(qn('w:fill'), 'E8EEF2')
                cell._tc.get_or_add_tcPr().append(shade)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def add_picture(paragraph, path: Path, caption: str, width: float, max_height: float) -> None:
    with Image.open(path) as image:
        image_width, image_height = image.size
    width = min(width, max_height * image_width / image_height)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_after = Pt(4)
    # python-docx embeds these exact bytes. Only the Word display size changes.
    shape = paragraph.add_run().add_picture(str(path), width=Inches(width))
    shape._inline.docPr.set('descr', plain(caption))


def add_pair(document, pair: list[tuple[str, Path]]) -> None:
    names = [path.stem for _, path in pair]
    if names[0].endswith('-light') and names[1] == names[0].removesuffix('-light') + '-dark':
        with Image.open(pair[0][1]) as first, Image.open(pair[1][1]) as second:
            if first.size != second.size:
                raise ValueError('Light and dark comparison images must have the same pixel dimensions.')
    table = document.add_table(rows=1, cols=2)
    table_style(table, [BODY_WIDTH / 2, BODY_WIDTH / 2])
    for cell, (caption, path) in zip(table.rows[0].cells, pair):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        add_picture(cell.paragraphs[0], path, caption, width=3.1, max_height=7.1)
        paragraph = cell.add_paragraph(style='Caption')
        paragraph.paragraph_format.space_after = Pt(3)
        add_inline(paragraph, caption)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def add_source(document, source: Path, source_index: int) -> None:
    previous_pair = False
    first_heading = True
    for kind, value in blocks(source):
        if kind == 'heading':
            depth, title = value
            if source_index == 0 and first_heading:
                paragraph = document.add_paragraph(shared.heading_text(title), style='Title')
            else:
                level = min(3, max(1, depth - (1 if source_index == 0 else 0)))
                paragraph = document.add_heading(shared.heading_text(title), level=level)
                paragraph.paragraph_format.page_break_before = (
                    first_heading or (source_index == 0 and bool(re.match(r'^(?:2|[4-8])\. ', title)))
                )
            first_heading = False
        elif kind == 'pair':
            if previous_pair:
                document.add_page_break()
            add_pair(document, value)
        elif kind == 'image':
            caption, path = value
            add_picture(document.add_paragraph(), path, caption, BODY_WIDTH, 7.1)
            add_inline(document.add_paragraph(style='Caption'), caption)
        elif kind == 'table':
            add_table(document, value)
        elif kind == 'code':
            paragraph = document.add_paragraph()
            run = paragraph.add_run(value)
            run.font.name = 'DejaVu Sans Mono'
            run.font.size = Pt(9)
        else:
            bullet = kind == 'list' and value.startswith(('- ', '* '))
            paragraph = document.add_paragraph(style='List Bullet' if bullet else 'Normal')
            if kind == 'list' and not bullet:
                paragraph.paragraph_format.left_indent = Inches(.2)
                paragraph.paragraph_format.first_line_indent = Inches(-.2)
            add_inline(paragraph, value[2:] if bullet else value)
        previous_pair = kind == 'pair'


def required_text() -> list[str]:
    """Check every source block, including all case fields, not a token sample."""
    required = []
    for source in SOURCES:
        for kind, value in blocks(source):
            if kind == 'heading':
                required.append(shared.heading_text(value[1]))
            elif kind == 'pair':
                required.extend(plain(caption) for caption, _ in value)
            elif kind == 'image':
                required.append(plain(value[0]))
            elif kind == 'table':
                required.extend(plain(cell) for row in value for cell in row)
            else:
                if kind == 'list' and value.startswith(('- ', '* ')):
                    value = value[2:]
                required.append(plain(value))
    return required


def check_document() -> None:
    with zipfile.ZipFile(OUTPUT) as archive:
        if archive.testzip() is not None:
            raise ValueError('The Word archive is damaged.')
        document = ElementTree.fromstring(archive.read('word/document.xml'))
        paragraphs = [''.join(p.itertext()) for p in document.findall('.//w:p', NS)]
        text = normalized(' '.join(paragraphs))
        for value in required_text():
            if normalized(value) not in text:
                raise ValueError(f'Missing editable Word content: {value[:150]}')
        if not document.findall('.//w:pStyle[@w:val="Title"]', NS):
            raise ValueError('The Word document needs an editable Title paragraph.')
        if len(document.findall('.//w:drawing', NS)) != len(image_uses()):
            raise ValueError('Every referenced screenshot must appear in the Word document.')
        embedded = {
            hashlib.sha256(archive.read(name)).hexdigest()
            for name in archive.namelist() if name.startswith('word/media/')
        }
        originals = {digest(path) for _, path in image_uses()}
        if embedded != originals:
            raise ValueError('Embedded images must be the exact original screenshot bytes.')
        styles = ElementTree.fromstring(archive.read('word/styles.xml'))
        title_styles = styles.findall('.//w:style[@w:styleId="Title"]', NS)
        if any(style.findall('.//w:pBdr', NS) for style in title_styles):
            raise ValueError('The Title style must not contain a decorative border.')
        for paragraph in document.findall('.//w:p', NS):
            if paragraph.find('./w:pPr/w:pStyle[@w:val="Title"]', NS) is not None:
                if paragraph.find('./w:pPr/w:pBdr', NS) is not None:
                    raise ValueError('The title paragraph must not contain a decorative border.')


def current_record() -> dict:
    record = json.loads(RECORD.read_text(encoding='utf-8'))
    if record.get('inputs') != inputs():
        raise ValueError('Team pack inputs changed. Rebuild, render and review every page.')
    if record.get('docx') != str(OUTPUT.relative_to(ROOT)) or record.get('docx_sha256') != digest(OUTPUT):
        raise ValueError('The Word document does not match its build record.')
    check_document()
    return record


def verify() -> None:
    record = current_record()
    review = record.get('visual_review', {})
    if (review.get('status') != 'passed' or not isinstance(review.get('pages'), int)
            or review['pages'] < 1 or review.get('docx_sha256') != digest(OUTPUT)):
        raise ValueError('Render and inspect every page, then run --record-review N.')
    print(f'Team pack inputs, {len(image_uses())} original screenshot uses and reviewed DOCX match.')


def record_review(pages: int) -> None:
    if pages < 1:
        raise ValueError('The reviewed page count must be positive.')
    record = current_record()
    record['visual_review'] = {'status': 'passed', 'pages': pages, 'docx_sha256': digest(OUTPUT)}
    RECORD.write_text(json.dumps(record, indent=2) + '\n', encoding='utf-8')
    print(f'Recorded human or agent visual inspection of all {pages} rendered pages.')


def build() -> None:
    before = inputs()
    document = Document()
    configure(document)
    for index, source in enumerate(SOURCES):
        add_source(document, source, index)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    check_document()
    if before != inputs():
        raise ValueError('Inputs changed during authoring. Rebuild from a stable set of sources.')
    record = {
        'format': 'docx', 'sources': [str(source.relative_to(ROOT)) for source in SOURCES],
        'docx': str(OUTPUT.relative_to(ROOT)), 'docx_sha256': digest(OUTPUT),
        'screenshot_uses': len(image_uses()),
        'visual_review': {'status': 'pending', 'pages': None, 'docx_sha256': None},
        'inputs': before,
    }
    RECORD.write_text(json.dumps(record, indent=2) + '\n', encoding='utf-8')
    print(f'Built editable Word team pack: {OUTPUT}')
    print('Visual review is pending. Render and inspect every page before recording review.')


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    mode = parser.add_mutually_exclusive_group()
    mode.add_argument('--check', action='store_true', help='Check freshness, content and recorded visual review')
    mode.add_argument('--record-review', type=int, metavar='PAGES', help='Record actual inspection of every rendered page')
    arguments = parser.parse_args()
    try:
        if arguments.check:
            verify()
        elif arguments.record_review is not None:
            record_review(arguments.record_review)
        else:
            build()
    except (ValueError, FileNotFoundError, zipfile.BadZipFile) as error:
        raise SystemExit(str(error)) from error
