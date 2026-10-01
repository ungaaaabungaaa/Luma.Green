#!/usr/bin/env python3
"""Build the editable Word guide from Markdown and unchanged browser captures.

Install scripts/user-guide-requirements.txt. Run from the repository root.
After rendering and inspecting every page, record review with --record-review N.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'docs/user-guide/guide.md'
OUTPUT = ROOT / 'output/docx/luma-green-user-guide.docx'
LOCK = ROOT / 'docs/user-guide/build.json'
BODY_WIDTH = 6.85


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def image_paths() -> list[Path]:
    return [(SOURCE.parent / item).resolve() for item in re.findall(r'!\[[^\]]*\]\(([^)]+)\)', SOURCE.read_text())]


def input_paths() -> list[Path]:
    paths = [SOURCE, Path(__file__).resolve(), ROOT / 'scripts/user-guide-requirements.txt']
    paths += image_paths()
    paths += sorted(SOURCE.parent.rglob('*captures.json'))
    return list(dict.fromkeys(paths))


def inputs() -> dict[str, str]:
    return {str(path.relative_to(ROOT)): digest(path) for path in input_paths()}


def check_document() -> None:
    with zipfile.ZipFile(OUTPUT) as archive:
        document = ElementTree.fromstring(archive.read('word/document.xml'))
        text = ' '.join(document.itertext())
        for required in ['Kabadiwala', 'Admin', 'Synthetic documentation fixture', 'not a GST', 'Evidence and maintenance']:
            if required not in text:
                raise SystemExit(f'Missing required Word content: {required}')
        namespace = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        if not document.findall('.//w:pStyle[@w:val="Title"]', namespace):
            raise SystemExit('The Word guide requires an editable Title style.')
        if len(document.findall('.//w:drawing', namespace)) != len(image_paths()):
            raise SystemExit('The Word guide must include every screenshot use.')
        if archive.testzip() is not None:
            raise SystemExit('The Word document archive is damaged.')


def verify_inputs(record: dict) -> None:
    if record['inputs'] != inputs():
        raise SystemExit('Guide inputs changed. Review, rebuild and inspect the Word guide.')


def verify() -> None:
    record = json.loads(LOCK.read_text())
    verify_inputs(record)
    check_document()
    if record['docx'] != str(OUTPUT.relative_to(ROOT)) or digest(OUTPUT) != record['docx_sha256']:
        raise SystemExit('Word guide does not match its build record.')
    review = record['visual_review']
    if review['status'] != 'passed' or review['docx_sha256'] != digest(OUTPUT) or review['pages'] < 1:
        raise SystemExit('Render and inspect all pages, then record review with --record-review N.')
    print(f'Guide source, {len(image_paths())} image uses and reviewed Word guide match the build record.')


def record_review(pages: int) -> None:
    if pages < 1:
        raise SystemExit('The reviewed page count must be positive.')
    record = json.loads(LOCK.read_text())
    verify_inputs(record)
    check_document()
    if digest(OUTPUT) != record['docx_sha256']:
        raise SystemExit('The document changed after building. Rebuild and inspect before recording review.')
    record['visual_review'] = {'status': 'passed', 'pages': pages, 'docx_sha256': digest(OUTPUT)}
    LOCK.write_text(json.dumps(record, indent=2) + '\n')
    print(f'Recorded inspection of all {pages} rendered pages.')


def inline(paragraph, value: str, bold: bool = False) -> None:
    """Keep ordinary text, emphasis, inline code and links editable."""
    pattern = r'(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))'
    for token in re.split(pattern, value):
        if not token:
            continue
        if token.startswith('**') and token.endswith('**'):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith('`') and token.endswith('`'):
            run = paragraph.add_run(token[1:-1])
            run.font.name = 'DejaVu Sans Mono'
            run.font.size = Pt(9)
        elif match := re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)', token):
            label, url = match.groups()
            # Visible URLs survive Word and Google Docs import without hidden destinations.
            run = paragraph.add_run(f'{label} ({url})' if label != url else label)
        else:
            run = paragraph.add_run(token)
        if bold:
            run.bold = True


def heading_text(value: str) -> str:
    # Simple heading wording makes Word navigation and Docs imports easier to scan.
    return re.sub(r'\s+', ' ', re.sub(r'[^\w\s]', ' ', value)).strip()


def configure(document) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(.65)
    section.left_margin = section.right_margin = Inches(.71)
    for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'Caption', 'List Bullet', 'List Number']:
        style = document.styles[name]
        style.font.name = 'DejaVu Sans'
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.underline = None
        style.paragraph_format.widow_control = True
        for border in list(style.element.xpath('./w:pPr/w:pBdr')):
            border.getparent().remove(border)
    normal = document.styles['Normal']
    normal.font.size = Pt(10)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.space_after = Pt(7)
    for name, size in [('Title', 28), ('Heading 1', 20), ('Heading 2', 12), ('Heading 3', 11)]:
        style = document.styles[name]
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(9)
        style.paragraph_format.space_after = Pt(9)
        style.paragraph_format.keep_with_next = True
    document.styles['Caption'].font.size = Pt(8)
    document.styles['Caption'].paragraph_format.space_after = Pt(9)
    # Page numbers help readers use this long operational manual.
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    run._r.addnext(field)
    document.core_properties.title = 'Luma Green platform user guide'
    document.core_properties.author = 'Luma.Green'
    document.core_properties.subject = 'Platform workflows and browser screenshot evidence'


def add_table(document, rows: list[list[str]]) -> None:
    count = len(rows[0])
    if any(len(row) != count for row in rows):
        raise ValueError('A guide table has inconsistent column counts.')
    proportions = [.29, .71] if count == 2 else [.23, .30, .47] if count == 3 else [1 / count] * count
    table = document.add_table(rows=0, cols=count)
    table.autofit = False
    for column, proportion in zip(table.columns, proportions):
        column.width = Inches(BODY_WIDTH * proportion)
    props = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement(f'w:{edge}')
        for name, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
            border.set(qn(f'w:{name}'), value)
        borders.append(border)
    props.append(borders)
    margins = OxmlElement('w:tblCellMar')
    for edge in ['top', 'left', 'bottom', 'right']:
        margin = OxmlElement(f'w:{edge}')
        margin.set(qn('w:w'), '100')
        margin.set(qn('w:type'), 'dxa')
        margins.append(margin)
    props.append(margins)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        if row_index == 0:
            row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for cell, value, proportion in zip(row.cells, values, proportions):
            cell.width = Inches(BODY_WIDTH * proportion)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(1)
            paragraph.paragraph_format.line_spacing = 1.05
            inline(paragraph, value, bold=row_index == 0)
            for run in paragraph.runs:
                run.font.size = Pt(8.5)
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), 'E8EEF2' if row_index == 0 else 'FFFFFF')
            cell._tc.get_or_add_tcPr().append(shade)
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def build() -> None:
    source = SOURCE.read_text()
    metadata = re.search(r'Edition: ([^.]+)\. Source baseline: ([a-f0-9]+)\.', source)
    if metadata is None:
        raise ValueError('The guide must declare its edition date and source baseline.')
    edition, checkpoint = metadata.groups()
    document = Document()
    configure(document)
    lines = source.splitlines()
    index = 0
    first_break = True
    next_page = False
    while index < len(lines):
        line = lines[index].strip()
        index += 1
        if not line:
            continue
        if line == '---':
            next_page = True
            if first_break:
                document.add_heading('Contents', level=1).paragraph_format.page_break_before = True
                for item in re.findall(r'^## (.+)$', source, re.MULTILINE):
                    paragraph = document.add_paragraph(heading_text(item))
                    paragraph.paragraph_format.space_after = Pt(2)
                    paragraph.paragraph_format.line_spacing = 1
                    for run in paragraph.runs:
                        run.font.size = Pt(9)
                first_break = False
            continue
        if match := re.fullmatch(r'!\[([^\]]*)\]\(([^)]+)\)', line):
            caption, relative = match.groups()
            path = (SOURCE.parent / relative).resolve()
            with Image.open(path) as picture:
                width, height = picture.size
            max_height = 6.5 if height > width * 1.3 else 4.4
            # These reference captures need a different balance of image and prose.
            if path.name == 'household-tracking.png':
                max_height = 8.4
            elif path.name == 'manufacturer-compliance.png':
                max_height = 8.0
            elif path.name == 'yard-trades.png':
                max_height = 3.8
            # Keep small mobile captures sharp instead of enlarging them to desktop width.
            image_width = min(BODY_WIDTH, max_height * width / height, width / 110)
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.page_break_before = next_page
            next_page = False
            if image_width < 3.8:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(4)
            shape = paragraph.add_run().add_picture(str(path), width=Inches(image_width))
            shape._inline.docPr.set('descr', caption)
            inline(document.add_paragraph(style='Caption'), caption)
            if path.name == 'household-tracking.png':
                next_page = True
            continue
        if line.startswith('|'):
            rows = [line]
            while index < len(lines) and lines[index].strip().startswith('|'):
                rows.append(lines[index].strip())
                index += 1
            values = []
            for row in rows:
                cells = [cell.strip() for cell in row.strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?', cell) for cell in cells):
                    values.append(cells)
            add_table(document, values)
            continue
        if match := re.match(r'^(#{1,4}) (.+)$', line):
            level, title = match.groups()
            paragraph = document.add_paragraph(heading_text(title), style='Title' if len(level) == 1 else f'Heading {len(level) - 1}')
            paragraph.paragraph_format.page_break_before = next_page or title == 'Error reports and owner settings'
            next_page = False
        elif re.match(r'^(?:\d+\.|-) ', line):
            # Preserve procedure step numbers exactly; each procedure restarts in source.
            paragraph = document.add_paragraph(style='List Bullet' if line.startswith('- ') else 'Normal')
            if line.startswith('- '):
                line = line[2:]
            else:
                paragraph.paragraph_format.left_indent = Inches(.17)
                paragraph.paragraph_format.first_line_indent = Inches(-.17)
            inline(paragraph, line)
        else:
            parts = [line]
            while index < len(lines) and lines[index].strip() and not re.match(r'^(?:#|\||!\[|---|\d+\. |- )', lines[index].strip()):
                parts.append(lines[index].strip())
                index += 1
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.page_break_before = next_page
            next_page = False
            inline(paragraph, ' '.join(parts))
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    check_document()
    record = {
        'format': 'docx',
        'edition': edition,
        'application_checkpoint': checkpoint,
        'docx': str(OUTPUT.relative_to(ROOT)),
        'docx_sha256': digest(OUTPUT),
        'visual_review': {'status': 'pending', 'pages': None, 'docx_sha256': None},
        'inputs': inputs(),
    }
    LOCK.write_text(json.dumps(record, indent=2) + '\n')
    print(f'Built editable Word guide with {len(image_paths())} image uses: {OUTPUT}')
    print('Render and inspect every page before recording visual review.')


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    mode = parser.add_mutually_exclusive_group()
    mode.add_argument('--check', action='store_true', help='Verify inputs, DOCX and recorded visual review')
    mode.add_argument('--record-review', type=int, metavar='PAGES', help='Record completed inspection of every rendered page')
    args = parser.parse_args()
    if args.check:
        verify()
    elif args.record_review is not None:
        record_review(args.record_review)
    else:
        build()
