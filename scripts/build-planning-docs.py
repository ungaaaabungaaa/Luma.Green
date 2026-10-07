#!/usr/bin/env python3
"""Build the two maintained planning documents with python-docx.

Derived from the reviewed interview/launch builders used on 7 October 2026.
No cloud writes, network calls, secrets, screenshots or render attestation.
"""
from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.text.hyperlink import Hyperlink

from document_links import add_document_hyperlink, resolve_document_link


def set_font(style, name='Arial', size=9.2, bold=False):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:ascii'), name)
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:hAnsi'), name)


def hyperlink(paragraph, label, url, link_base):
    source, repository = link_base
    destination = resolve_document_link(url, source, repository)
    add_document_hyperlink(paragraph, label, destination)


TOKEN = re.compile(r'\[([^\]]+)\]\(([^)]+)\)|\*\*([^*]+)\*\*|`([^`]+)`|_([^_]+)_')


def add_inline(paragraph, value, link_base):
    value = value.replace('  ', ' ')
    cursor = 0
    for match in TOKEN.finditer(value):
        if match.start() > cursor:
            paragraph.add_run(value[cursor:match.start()])
        if match.group(1) is not None:
            hyperlink(paragraph, match.group(1), match.group(2), link_base)
        elif match.group(3) is not None:
            paragraph.add_run(match.group(3)).bold = True
        elif match.group(4) is not None:
            run = paragraph.add_run(match.group(4))
            run.font.name = 'Liberation Mono'
            run.font.size = Pt(8.8)
        elif match.group(5) is not None:
            paragraph.add_run(match.group(5)).italic = True
        cursor = match.end()
    if cursor < len(value):
        paragraph.add_run(value[cursor:])


def clean_heading(value):
    value = value.replace('Luma.Green', 'Luma Green')
    value = value.replace('→', ' to ').replace('/', ' or ')
    value = re.sub(r'[^\w\s]', ' ', value, flags=re.UNICODE)
    return re.sub(r'\s+', ' ', value).strip()


def parse_table_row(value):
    return [cell.strip() for cell in value.strip().strip('|').split('|')]


def style_launch_document(doc):
    # Use consistent bordered tables with explicit widths and readable header rows.
    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        widths = (1.78, 5.28)
        for col, width in zip(table.columns, widths):
            col.width = Inches(width)
        props = table._tbl.tblPr
        borders = OxmlElement('w:tblBorders')
        for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            el = OxmlElement('w:' + edge)
            for key, val in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
                el.set(qn('w:' + key), val)
            borders.append(el)
        props.append(borders)
        for idx, row in enumerate(table.rows):
            trpr = row._tr.get_or_add_trPr()
            trpr.append(OxmlElement('w:cantSplit'))
            if idx == 0:
                trpr.append(OxmlElement('w:tblHeader'))
            for ci, cell in enumerate(row.cells):
                cell.width = Inches(widths[ci])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tcpr = cell._tc.get_or_add_tcPr()
                shd = OxmlElement('w:shd')
                shd.set(qn('w:fill'), 'DCE6EE' if idx == 0 else 'FFFFFF')
                tcpr.append(shd)
                margins = OxmlElement('w:tcMar')
                for edge in ('top', 'left', 'bottom', 'right'):
                    value = OxmlElement('w:' + edge)
                    value.set(qn('w:w'), '85')
                    value.set(qn('w:type'), 'dxa')
                    margins.append(value)
                tcpr.append(margins)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.line_spacing = 1.02
                    inline_runs = [run for part in p.iter_inner_content() for run in (part.runs if isinstance(part, Hyperlink) else [part])]
                    for run in inline_runs:
                        run.font.name = 'Arial'
                        run.font.size = Pt(9.5)
                        run.bold = idx == 0
                        run.font.color.rgb = RGBColor(0, 0, 0)
    for paragraph in doc.paragraphs:
        previous = paragraph._p.getprevious()
        if previous is not None and previous.tag == qn('w:tbl'):
            paragraph.paragraph_format.space_before = Pt(7)
    # Page numbers help the team cite test evidence in this long plan.
    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = 2
    run = footer.add_run('Luma Green Launch Test Plan  |  ')
    run.font.name = 'Arial'
    run.font.size = Pt(8)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)


def build_document(repository: Path, kind: str) -> Path:
    launch = kind == "launch"
    source = repository / ("docs/testing/launch-2026-10-10.md" if launch else "docs/product/platform-refinement-interview.md")
    output = repository / "output/docx" / ("luma-green-launch-test-plan.docx" if launch else "luma-green-platform-refinement-interview.docx")
    title = "Luma Green Launch Test Plan" if launch else "Luma Green Platform Refinement Interview"
    subject = "Launch acceptance plan for 10 October 2026" if launch else "Running product requirements interview"
    link_base = (source, repository)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5 if launch else 8.27)
    section.page_height = Inches(11 if launch else 11.69)
    section.top_margin = Inches(0.7 if launch else 0.66)
    section.bottom_margin = Inches(0.7 if launch else 0.66)
    section.left_margin = Inches(0.72 if launch else 0.73)
    section.right_margin = Inches(0.72 if launch else 0.73)

    styles = doc.styles
    normal = styles['Normal']
    set_font(normal, size=11 if launch else 9.2)
    normal.paragraph_format.space_after = Pt(5 if launch else 3)
    normal.paragraph_format.line_spacing = 1.04 if launch else 1.08
    normal.paragraph_format.keep_together = True
    for name, size, before, after in [
        ('Title', 19, 0, 9),
        ('Heading 1', 14, 11, 5),
        ('Heading 2', 11, 8, 3),
        ('Heading 3', 10, 6, 2),
    ]:
        style = styles[name]
        set_font(style, size=size, bold=True)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    title_ppr = styles['Title']._element.get_or_add_pPr()
    for border in list(title_ppr.findall(qn('w:pBdr'))):
        title_ppr.remove(border)
    for name in ('List Bullet', 'List Number'):
        style = styles[name]
        set_font(style, size=11 if launch else 9.2)
        style.paragraph_format.space_after = Pt(2)
        style.paragraph_format.line_spacing = 1.04 if launch else 1.08

    code_style = styles.add_style('Luma Code', 1)
    set_font(code_style, name='Liberation Mono', size=8.2)
    code_style.paragraph_format.left_indent = Inches(0.16)
    code_style.paragraph_format.space_after = Pt(1)
    code_style.paragraph_format.line_spacing = 1.02

    lines = source.read_text(encoding='utf-8').splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            i += 1
            continue
        if line.startswith('```'):
            i += 1
            while i < len(lines) and not lines[i].startswith('```'):
                paragraph = doc.add_paragraph(style='Luma Code')
                paragraph.add_run(lines[i])
                i += 1
            i += 1
            continue
        if line.startswith('# '):
            doc.add_paragraph(title, style='Title')
            i += 1
            continue
        if line.startswith('## '):
            doc.add_paragraph(clean_heading(line[3:]), style='Heading 1')
            i += 1
            continue
        if line.startswith('#### '):
            doc.add_paragraph(clean_heading(line[5:]), style='Heading 3')
            i += 1
            continue
        if line.startswith('### '):
            doc.add_paragraph(clean_heading(line[4:]), style='Heading 2')
            i += 1
            continue
        if line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].startswith('|'):
                rows.append(parse_table_row(lines[i]))
                i += 1
            if len(rows) < 3:
                continue
            headers = rows[0]
            data = rows[2:]
            if len(headers) <= 2:
                table = doc.add_table(rows=1, cols=len(headers))
                table.style = 'Table Grid' if launch else 'Light Shading Accent 1'
                for col, label in enumerate(headers):
                    table.rows[0].cells[col].text = label
                for row in data:
                    cells = table.add_row().cells
                    for col, cell in enumerate(row):
                        if col < len(cells):
                            cells[col].text = cell
            else:
                for row in data:
                    if not row or not row[0]:
                        continue
                    title = doc.add_paragraph(clean_heading(row[0]), style='Heading 3')
                    title.paragraph_format.keep_with_next = True
                    for col, cell in enumerate(row[1:], 1):
                        if col >= len(headers):
                            break
                        paragraph = doc.add_paragraph()
                        paragraph.paragraph_format.left_indent = Inches(0.13)
                        paragraph.paragraph_format.keep_with_next = col < len(row) - 1
                        paragraph.add_run(headers[col] + ': ').bold = True
                        add_inline(paragraph, cell, link_base)
            continue
        if line.startswith('> '):
            content = []
            while i < len(lines) and lines[i].startswith('> '):
                content.append(lines[i][2:].strip())
                i += 1
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.14)
            paragraph.paragraph_format.space_after = Pt(8)
            add_inline(paragraph, ' '.join(content), link_base)
            continue
        list_match = re.match(r'^(\s*)([-*]|\d+\.)\s+(.*)', line)
        if list_match:
            indent, marker, item = list_match.groups()
            i += 1
            while i < len(lines) and lines[i].startswith('  ') and lines[i].strip() and not re.match(r'^\s*([-*]|\d+\.)\s+', lines[i]):
                item += ' ' + lines[i].strip()
                i += 1
            is_number = marker.endswith('.') and marker[:-1].isdigit()
            paragraph = doc.add_paragraph(style='Normal' if is_number else 'List Bullet')
            if is_number:
                paragraph.paragraph_format.left_indent = Inches(0.30)
                paragraph.paragraph_format.first_line_indent = Inches(-0.22)
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.add_run(marker + '  ')
            elif indent:
                paragraph.paragraph_format.left_indent = Inches(0.52)
            add_inline(paragraph, item, link_base)
            continue
        content = [line.strip()]
        i += 1
        while i < len(lines):
            next_line = lines[i]
            if not next_line.strip() or re.match(r'^(#{1,4} |```|> |\||\s*[-*] |\s*\d+\. )', next_line):
                break
            content.append(next_line.strip())
            i += 1
        if content[0].startswith('The editable Word copy is'):
            continue
        paragraph = doc.add_paragraph()
        add_inline(paragraph, ' '.join(content), link_base)

    # Keep comparison rows intact and repeat labels after a page break.
    for table in doc.tables:
        for index, row in enumerate(table.rows):
            properties = row._tr.get_or_add_trPr()
            properties.append(OxmlElement('w:cantSplit'))
            if index == 0:
                properties.append(OxmlElement('w:tblHeader'))

    if launch:
        style_launch_document(doc)
    doc.core_properties.title = title
    doc.core_properties.subject = subject
    doc.core_properties.author = 'Luma.Green'
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    return output


def main():
    parser = argparse.ArgumentParser(description="Rebuild the maintained Luma interview and launch Word documents. Rendering and visual review are separate required steps.")
    parser.add_argument("--document", choices=("interview", "launch", "all"), default="all")
    parser.add_argument("--root", type=Path, default=Path(__file__).resolve().parents[1], help="Repository root; defaults to this script's parent repository.")
    args = parser.parse_args()
    repository = args.root.resolve()
    kinds = ("interview", "launch") if args.document == "all" else (args.document,)
    for kind in kinds:
        print(build_document(repository, kind))


if __name__ == "__main__":
    main()
