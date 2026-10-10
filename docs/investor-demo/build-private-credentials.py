#!/usr/bin/env python3
"""Project only name/email/password from private input into a private Word annex."""
import argparse
import json
import os
import re
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt


def private_path(raw: str, suffix: str) -> Path:
    path = Path(raw).resolve()
    if '.convex' not in path.parts or path.suffix.lower() != suffix:
        raise ValueError('Input and output must be private .convex paths with the required extension.')
    return path


def build(source: Path, output: Path, expected: int) -> None:
    data = json.loads(source.read_text())
    rows = data['accounts'] if isinstance(data, dict) else data
    if not isinstance(rows, list) or len(rows) != expected:
        raise ValueError('Private account count does not match the requested environment roster.')
    safe = []
    seen = set()
    for row in rows:
        values = {key: row.get(key) for key in ['name', 'email', 'password']}
        if any(not isinstance(value, str) or not value.strip() for value in values.values()):
            raise ValueError('Each private account needs name, email and password strings.')
        if not re.fullmatch(r'[a-z0-9-]+-[1-5]@investor\.luma\.invalid', values['email']):
            raise ValueError('Only the agreed fictional investor account namespace is accepted.')
        if values['email'] in seen:
            raise ValueError('The private account list contains a duplicate.')
        seen.add(values['email'])
        safe.append(values)
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.65)
    section.left_margin = section.right_margin = Inches(.7)
    document.styles['Normal'].font.name = 'DejaVu Sans'
    document.styles['Normal'].font.size = Pt(11)
    document.core_properties.title = 'Restricted investor demonstration accounts'
    document.core_properties.author = 'Luma.Green'
    for index, row in enumerate(safe):
        if index and index % 5 == 0:
            document.add_page_break()
        for key in ['name', 'email', 'password']:
            p = document.add_paragraph()
            p.add_run(key.capitalize() + ': ').bold = True
            p.add_run(row[key])
        document.add_paragraph()
    output.parent.mkdir(parents=True, exist_ok=True)
    old_mask = os.umask(0o077)
    try:
        document.save(output)
        output.chmod(0o600)
    finally:
        os.umask(old_mask)
    print('Private credentials Word file written. No account values were logged.')


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input', required=True)
    parser.add_argument('--output', required=True)
    parser.add_argument('--expected-count', type=int, default=140)
    args = parser.parse_args()
    build(private_path(args.input, '.json'), private_path(args.output, '.docx'), args.expected_count)
